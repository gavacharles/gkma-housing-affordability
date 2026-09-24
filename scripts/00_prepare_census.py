"""Census 2024 parish profiles (UBOS NPHC 2024 sub-county profile tables).

Source: data/external/ubos/harvest/NPHC-2024-Subcounty-Profiles-Excel-Tables.xlsx,
downloaded from https://www.ubos.org/datasets/ with tools/cio_pipeline-2's
harvester. Each sheet stacks district / county / sub-county / parish rows in
column A; the level is given by the cell indent (0, 1, 2, 3).

Writes
  data/external/census2024_parish.csv   one row per GKMA parish: households,
      household population and size, and shares of households with TV,
      computer, grid electricity, improved water, improved sanitation, and in
      the subsistence economy, plus a wealth index (first principal component)
  data/external/census2024_parish.gpkg  the same joined to the UBOS 2016 parish
      polygons (name match within sub-county; unmatched parishes are listed)
"""
import re

import _common  # noqa: F401
import geopandas as gpd
import numpy as np
import openpyxl
import pandas as pd
from rapidfuzz import fuzz, process

from gkma.config import load_config, p

SRC = p("data/external/ubos/harvest/NPHC-2024-Subcounty-Profiles-Excel-Tables.xlsx")
GKMA = {"KAMPALA", "WAKISO", "MUKONO", "MPIGI", "BUIKWE", "LUWERO", "LUWEERO"}
TABLES = {  # sheet -> {output column: header column index}
    "Table2": {"household_pop": 1, "households": 2, "hh_size": 3},
    "Table7": {"n_radio": 1, "n_tv": 2, "n_computer": 3},
    "Table12": {"n_unimproved_water": 1, "n_improved_water": 2, "n_improved_sanitation": 3,
                "n_unimproved_sanitation": 4, "n_open_defecation": 5},
    "Table13": {"n_grid": 1, "n_solar": 2},
    "Table14": {"n_subsistence": 1, "n_pdm": 2},
}


def read_table(wb, sheet, cols):
    ws = wb[sheet]
    rows, path = [], {}
    for r in range(7, ws.max_row + 1):
        c = ws.cell(r, 1)
        if c.value is None or c.alignment.indent is None:
            continue
        level = int(c.alignment.indent)
        path[level] = str(c.value).strip()
        for deeper in [k for k in path if k > level]:
            del path[deeper]
        if level != 3 or path.get(0) not in GKMA:
            continue
        rec = {"district": path.get(0), "county": path.get(1), "subcounty": path.get(2), "parish": path[3]}
        for name, j in cols.items():
            rec[name] = ws.cell(r, j + 1).value
        rows.append(rec)
    return pd.DataFrame(rows)


wb = openpyxl.load_workbook(SRC, read_only=False, data_only=True)
keys = ["district", "county", "subcounty", "parish"]
df = None
for sheet, cols in TABLES.items():
    t = read_table(wb, sheet, cols)
    df = t if df is None else df.merge(t, on=keys, how="left")
for c in df.columns.difference(keys):
    df[c] = pd.to_numeric(df[c], errors="coerce")

df = df[df["households"] > 0].reset_index(drop=True)   # drop empty (e.g. institutional) parishes
df["district"] = df["district"].replace({"LUWEERO": "LUWERO"})
hh = df["households"]
df["sh_tv"] = df["n_tv"] / hh
df["sh_computer"] = df["n_computer"] / hh
df["sh_grid"] = df["n_grid"] / hh
df["sh_improved_water"] = df["n_improved_water"] / hh
df["sh_improved_sanitation"] = df["n_improved_sanitation"] / hh
df["sh_subsistence"] = df["n_subsistence"] / hh
wcols = ["sh_tv", "sh_computer", "sh_grid", "sh_improved_water", "sh_improved_sanitation", "sh_subsistence"]
Z = df[wcols].apply(lambda s: (s - s.mean()) / s.std()).fillna(0)
Z["sh_subsistence"] *= -1                       # subsistence lowers wealth
vals, vecs = np.linalg.eigh(np.cov(Z.T))
pc1 = vecs[:, -1] * np.sign(vecs[:, -1].sum())
df["wealth_index"] = Z.to_numpy() @ pc1
df["wealth_index"] = (df["wealth_index"] - df["wealth_index"].mean()) / df["wealth_index"].std()
explained = vals[-1] / vals.sum()
df.to_csv(p("data/external/census2024_parish.csv"), index=False)
print(f"{len(df)} GKMA parishes, {int(hh.sum()):,} households; wealth index PC1 explains {explained:.0%}")
print(df.groupby("district")[["households", "household_pop"]].sum().astype(int).to_string())

# Join to the 2016 parish polygons: match parish name within the same sub-county
g = load_config()["geography"]
par = gpd.read_file(p(g["parishes"]))


def key(s):
    s = re.sub(r"\b(division|town council|municipality|ward|sub ?county)\b", "", str(s).lower())
    return re.sub(r"[^a-z0-9]", "", s)


par["_d"], par["_s"], par["_p"] = par[g["district_name_col"]].map(key), par[g["subcounty_name_col"]].map(key), par[g["parish_name_col"]].map(key)
df["_d"], df["_s"], df["_p"] = df["district"].map(key), df["subcounty"].map(key), df["parish"].map(key)
match = []
for i, r in df.iterrows():
    cand = par[par["_d"] == r["_d"]]
    same_sub = cand[cand["_s"].map(lambda s: fuzz.ratio(s, r["_s"]) >= 85)]
    pool = same_sub if len(same_sub) else cand
    hit = process.extractOne(r["_p"], pool["_p"].tolist(), scorer=fuzz.ratio) if len(pool) else None
    match.append(pool.index[hit[2]] if hit and hit[1] >= 88 else None)
df["poly_idx"] = match
joined = par.join(df.dropna(subset=["poly_idx"]).drop_duplicates("poly_idx").set_index("poly_idx")
                  [["subcounty", "parish", "households", "household_pop", "hh_size", "wealth_index"] + wcols])
joined.to_file(p("data/external/census2024_parish.gpkg"), driver="GPKG")
print(f"matched {df['poly_idx'].notna().sum()} of {len(df)} census parishes to {joined['households'].notna().sum()} "
      f"of {len(par)} polygons ({joined['households'].sum() / hh.sum():.0%} of households)")
df[df["poly_idx"].isna()][["district", "subcounty", "parish", "households"]].to_csv(
    p("data/external/census2024_parish_unmatched.csv"), index=False)
