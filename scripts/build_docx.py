"""Convert the compiled Markdown manuscript and supplement to Word (.docx).

  python scripts/build_docx.py
  -> docs/manuscript/manuscript.docx, docs/manuscript/supplementary.docx

Handles the subset of Markdown the manuscript uses: headings, paragraphs with
**bold** / *italic*, bullet lists, numbered reference lists, pipe tables and
![caption](image) figures. Figures are inserted at 16 cm width with the
caption below; table captions (bold paragraphs starting "Table") stay above.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
MS = ROOT / "docs/manuscript"


def add_runs(par, text, size=None):
    """Add text with **bold** and *italic* spans (significance stars are kept literal)."""
    text = re.sub(r"(?<=[\d)])\*{1,3}|\*{1,3}(?= p [<=])", lambda m: "\x00" * len(m.group(0)), text)
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = par.add_run(tok[2:-2].replace("\x00", "*"))
            run.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = par.add_run(tok[1:-1].replace("\x00", "*"))
            run.italic = True
        else:
            run = par.add_run(tok.replace("\x00", "*"))
        if size:
            run.font.size = Pt(size)


def shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_fill)
    tc.append(sh)


def add_table(doc, lines):
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 8 if ncol <= 6 else 7
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, r[j] if j < len(r) else "", size=size)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                shade(cell, "E8EEF4")
    doc.add_paragraph()


def convert(src: Path, dst: Path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Palatino Linotype"
    st.font.size = Pt(10)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        l = lines[i]
        s = l.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        m = re.match(r"!\[(.*)\]\((.*)\)", s)
        if m:
            cap, path = m.groups()
            img = (src.parent / path).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"[Missing figure: {path}]")
            p = doc.add_paragraph()
            mm = re.match(r"((?:Figure|Table) S?\d+\.)(.*)", cap)
            if mm:
                r = p.add_run(mm.group(1))
                r.bold = True
                r.font.size = Pt(9)
                add_runs(p, mm.group(2), size=9)
            else:
                add_runs(p, cap, size=9)
            i += 1
            continue
        h = re.match(r"(#{1,4})\s+(.*)", s)
        if h:
            level = len(h.group(1))
            if level == 1:
                p = doc.add_paragraph()
                r = p.add_run(h.group(2))
                r.bold = True
                r.font.size = Pt(16)
            else:
                doc.add_heading(h.group(2), level=level - 1)
            i += 1
            continue
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, s[2:])
            i += 1
            continue
        if re.match(r"\d+\. ", s) and "References" in "".join(lines[max(0, i - 400):i]):
            p = doc.add_paragraph()
            add_runs(p, s, size=9)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            i += 1
            continue
        # paragraph: join soft-wrapped lines
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"(#|\||!\[|- |\d+\. )", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_runs(p, " ".join(buf))
        p.paragraph_format.space_after = Pt(6)
    doc.save(dst)
    print("wrote", dst.relative_to(ROOT))


convert(MS / "manuscript_draft.md", MS / "manuscript.docx")
convert(MS / "supplementary.md", MS / "supplementary.docx")
